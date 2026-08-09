"""Assemble the PPTX with live, editable text over rendered background plates."""
import json, pathlib, re
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent
PX   = 6350                      # 1920px -> 13.333in : EMU per CSS px
PT   = 0.5                       # 1920px -> 960pt    : pt per CSS px
ROUND= HERE/"round"; ROUND.mkdir(exist_ok=True)

ALIGN = {"start":PP_ALIGN.LEFT, "left":PP_ALIGN.LEFT, "center":PP_ALIGN.CENTER,
         "right":PP_ALIGN.RIGHT, "end":PP_ALIGN.RIGHT, "justify":PP_ALIGN.JUSTIFY}

NOTES = json.loads((HERE/"notes.json").read_text())

def rgb(css):
    m = re.findall(r"[\d.]+", css)
    return RGBColor(int(float(m[0])), int(float(m[1])), int(float(m[2])))

def family(css):
    return css.split(",")[0].strip().strip('"').strip("'")

def rounded(src, w, h, r=14):
    """Pre-round screenshot corners so they sit on the plate's rounded frame."""
    out = ROUND/(pathlib.Path(src).stem + f"_{w}x{h}.png")
    if out.exists(): return out
    im = Image.open(HERE/src).convert("RGBA").resize((w, h), Image.LANCZOS)
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, w-1, h-1], radius=r, fill=255)
    im.putalpha(mask); im.save(out)
    return out

def set_spacing(run, px):
    """CSS letter-spacing -> a:rPr@spc (hundredths of a point)."""
    if not px: return
    run.font._rPr.set("spc", str(int(round(px * PT * 100))))

data = json.loads((HERE/"layout.json").read_text())

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]

for i, s in enumerate(data, 1):
    sl = prs.slides.add_slide(blank)

    # 1. background plate (gradients, card fills, borders, rules, gradient numerals)
    sl.shapes.add_picture(str(HERE/"bg"/f"{s['id']}.png"), 0, 0,
                          width=prs.slide_width, height=prs.slide_height)

    # 2. screenshots as independent picture objects
    for im in s["images"]:
        w, h = int(round(im["w"])), int(round(im["h"]))
        sl.shapes.add_picture(str(rounded(im["src"], w*2, h*2)),
                              Emu(int(im["x"]*PX)), Emu(int(im["y"]*PX)),
                              width=Emu(int(im["w"]*PX)), height=Emu(int(im["h"]*PX)))

    # 3. live text
    for b in s["blocks"]:
        x = (b["x"] + b["padL"]) * PX
        y = (b["y"] + b["padT"]) * PX
        w = max(b["w"] - b["padL"] - b["padR"], 1) * PX
        h = max(b["h"] - b["padT"] - b["padB"], 1) * PX

        # Boxes are sized to Chromium's exact text width. PowerPoint's metrics differ
        # slightly, so an exact box wraps a word early ("assetdrills.co / m"). Single-line
        # blocks never wrap at all; multi-line blocks get a few px of slack.
        single_line = b["h"] <= b["lineHeight"] * 1.6
        if not single_line:
            w += 8 * PX

        tb = sl.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
        tf = tb.text_frame
        tf.word_wrap = not single_line
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.auto_size = None
        # never let PowerPoint shrink or grow the copy on its own
        bodyPr = tf._txBody.find(qn('a:bodyPr'))
        for tag in ('a:normAutofit', 'a:spAutoFit'):
            el = bodyPr.find(qn(tag))
            if el is not None: bodyPr.remove(el)

        p = tf.paragraphs[0]
        p.alignment = ALIGN.get(b["align"], PP_ALIGN.LEFT)
        p.line_spacing = Pt(b["lineHeight"] * PT)
        p.space_before = p.space_after = Pt(0)

        for r in b["runs"]:
            if r.get("br"):
                p.add_line_break()
                continue
            txt = r["text"]
            if b["transform"] == "uppercase": txt = txt.upper()
            run = p.add_run(); run.text = txt
            f = run.font
            f.name = family(r["family"] or b["family"])
            f.size = Pt(b["fontSize"] * PT)
            f.bold = r["bold"]; f.italic = r["italic"]
            f.color.rgb = rgb(r["color"])
            set_spacing(run, b["letterSpacing"])

    sl.notes_slide.notes_text_frame.text = NOTES.get(str(i), "")

dest = HERE.parent/"ASSET-Pit-Operations-Briefing-editable.pptx"
prs.save(dest)
print("saved:", dest.name, f"{dest.stat().st_size/1e6:.1f} MB",
      "| slides:", len(data),
      "| text boxes:", sum(len(s['blocks']) for s in data))
