import pathlib, glob, json
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Emu

HERE = pathlib.Path(__file__).parent
SRC  = sorted(glob.glob(str(HERE/"png"/"slide-*.png")))
TMP  = HERE/"pptx_img"; TMP.mkdir(exist_ok=True)

NOTES = json.loads((HERE/"notes.json").read_text())   # shared with build_editable.py

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

for i, f in enumerate(SRC, 1):
    im = Image.open(f)
    if im.width != 2560:
        im = im.resize((2560, 1440), Image.LANCZOS)
    out = TMP/f"s{i:02d}.png"
    im.save(out, optimize=True)

    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(out), 0, 0, width=prs.slide_width, height=prs.slide_height)
    s.notes_slide.notes_text_frame.text = NOTES.get(str(i), "")

dest = HERE.parent/"ASSET-Pit-Operations-Briefing.pptx"
prs.save(dest)
print("saved:", dest, f"{dest.stat().st_size/1e6:.1f} MB", len(SRC), "slides")
